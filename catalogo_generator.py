import os
import sys
import subprocess
from urllib.parse import quote

def asegurar_dependencias():
    try:
        import openpyxl
        import fpdf
    except ImportError:
        print("[!] Instalando dependencias...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "fpdf2"])

asegurar_dependencias()

import openpyxl
from fpdf import FPDF
from fpdf.enums import XPos, YPos


def carpeta_base():
    """Carpeta donde esta el .exe (o el script si no esta compilado)"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def cargar_config():
    cfg = {"NUMERO_WHATSAPP": "", "NOMBRE_NEGOCIO": "", "MONEDA": "$"}
    ruta = os.path.join(carpeta_base(), "config.txt")
    if os.path.exists(ruta):
        with open(ruta, encoding="utf-8") as f:
            for linea in f:
                if "=" in linea:
                    k, v = linea.split("=", 1)
                    cfg[k.strip()] = v.strip()
    return cfg


def crear_plantilla():
    ruta = os.path.join(carpeta_base(), "plantilla_ejemplo.xlsx")
    if os.path.exists(ruta):
        return ruta
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["Producto", "Precio", "Descripcion"])
    ws.append(["Playera Negra", "250", "Playera de algodon 100%, talla M"])
    ws.append(["Gorra Azul", "180", "Gorra ajustable con logo bordado"])
    ws.append(["Calcetines x3", "120", "Paquete de 3 pares"])
    wb.save(ruta)
    return ruta


def leer_excel(ruta):
    productos = []
    wb = openpyxl.load_workbook(ruta)
    ws = wb.active
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila and fila[0]:
            productos.append({
                "nombre": str(fila[0]),
                "precio": str(fila[1]) if fila[1] is not None else "0",
                "descripcion": str(fila[2]) if fila[2] is not None else "",
            })
    return productos


def link_whatsapp(producto, numero, moneda):
    msg = "Hola, me interesa: {} ({}{})".format(producto["nombre"], moneda, producto["precio"])
    return "https://wa.me/{}?text={}".format(numero, quote(msg))


def generar_pdf(productos, numero, negocio, moneda, salida):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    # ---- PORTADA ----
    pdf.add_page()
    pdf.set_fill_color(41, 128, 185)
    pdf.rect(0, 0, 210, 297, style="F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 34)
    pdf.cell(0, 40, negocio, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.set_font("Helvetica", "", 18)
    pdf.cell(0, 12, "Catalogo de Productos", new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.set_font("Helvetica", "", 13)
    pdf.cell(0, 10, "{} productos disponibles".format(len(productos)),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

    # ---- PRODUCTOS ----
    pdf.add_page()
    pdf.set_text_color(0, 0, 0)
    for i, p in enumerate(productos, 1):
        if pdf.get_y() > 235:
            pdf.add_page()
        y0 = pdf.get_y()
        if i % 2 == 0:
            pdf.set_fill_color(245, 245, 245)
            pdf.rect(10, y0, 190, 36, style="F")

        pdf.set_font("Helvetica", "B", 10)
        pdf.set_xy(14, y0 + 3)
        pdf.cell(10, 6, "#{}".format(i))

        pdf.set_font("Helvetica", "B", 13)
        pdf.set_xy(28, y0 + 2)
        pdf.cell(105, 7, p["nombre"][:45])

        pdf.set_text_color(39, 174, 96)
        pdf.set_xy(140, y0 + 2)
        pdf.cell(55, 7, moneda + p["precio"], align="R")
        pdf.set_text_color(0, 0, 0)

        pdf.set_font("Helvetica", "", 9)
        pdf.set_xy(28, y0 + 11)
        pdf.multi_cell(160, 5, p["descripcion"][:110])

        pdf.set_font("Helvetica", "U", 9)
        pdf.set_text_color(0, 102, 204)
        pdf.set_xy(28, pdf.get_y() + 1)
        pdf.cell(60, 6, "PEDIR POR WHATSAPP", link=link_whatsapp(p, numero, moneda))
        pdf.set_text_color(0, 0, 0)

        pdf.set_xy(10, y0 + 38)
        pdf.set_draw_color(200, 200, 200)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(4)

    pdf.output(salida)
    return salida


def main():
    cfg = cargar_config()
    print("=" * 50)
    print("   GENERADOR DE CATALOGO WHATSAPP")
    print("=" * 50)

    plantilla = crear_plantilla()

    print("\nRuta de tu Excel (Enter = usar plantilla_ejemplo.xlsx):")
    ruta = input(">>> ").strip().strip('"')
    if not ruta:
        ruta = plantilla
    if not os.path.exists(ruta):
        print("[ERROR] No existe el archivo: {}".format(ruta))
        return

    num = input("\nNumero de WhatsApp con codigo de pais (ej. 5215512345678) [{}]: ".format(cfg["NUMERO_WHATSAPP"]))
    numero = num.strip() or cfg["NUMERO_WHATSAPP"]

    neg = input("Nombre de tu negocio [{}]: ".format(cfg["NOMBRE_NEGOCIO"]))
    negocio = neg.strip() or cfg["NOMBRE_NEGOCIO"] or "Mi Negocio"

    print("\n[...] Generando catalogo...")
    productos = leer_excel(ruta)
    if not productos:
        print("[ERROR] El Excel no tiene productos (revisa que la fila 1 sean encabezados)")
        return

    salida = os.path.join(carpeta_base(), "catalogo_{}.pdf".format(negocio.replace(" ", "_")))
    generar_pdf(productos, numero, negocio, cfg["MONEDA"], salida)

    print("\n[OK] Catalogo listo con {} productos".format(len(productos)))
    print("[OK] Archivo: {}".format(salida))
    print("\nAbre el PDF y toca 'PEDIR POR WHATSAPP' para probar los links.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("[ERROR] {}".format(e))
    input("\nPresiona ENTER para cerrar...")
